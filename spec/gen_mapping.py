#!/usr/bin/env python3
"""Excel mapping spec -> Cloud Integration Groovy mapping step."""
import sys, re
from openpyxl import load_workbook

XLSX, OUT = sys.argv[1], sys.argv[2]
wb = load_workbook(XLSX)

def rows(sheet):
    ws = wb[sheet]; it = ws.iter_rows(values_only=True); hdr = next(it)
    return [dict(zip(hdr, r)) for r in it if r and r[0]]

fields = rows("FieldMapping")
rules  = sorted(rows("RiskRules"), key=lambda r: int(r["Priority"]))

def expr(path, transform):
    acc = "src" + "".join(f"?.get('{p}')" for p in path.split("."))
    t = (transform or "direct").strip()
    if t == "direct":                      return acc
    if t.startswith("ROUND"):
        d = re.search(r",\s*(\d+)\s*\)", t)
        return f"round1({acc})" if (d and d.group(1) == "1") else f"round1({acc})"
    if "UPPER" in t and "TRIM" in t:       return f"({acc} as String)?.trim()?.toUpperCase()"
    if t.startswith("DERIVED"):            return None
    return acc

def cond(c):
    c = (c or "").strip()
    if c.upper() == "DEFAULT": return "true"
    c = re.sub(r"\bAND\b", "&&", c, flags=re.I)
    c = re.sub(r"\bOR\b", "||", c, flags=re.I)
    for v in ("aqi", "tempF", "windMph", "pm25"):
        c = re.sub(rf"\b{v}\b", f"n(t.{v})", c)
    return c

L = []
L.append("import com.sap.gateway.ip.core.customdev.util.Message")
L.append("import groovy.json.JsonOutput")
L.append("import groovy.json.JsonSlurper")
L.append("")
L.append("// GENERATED from %s by Nova. Do not hand-edit; regenerate from the spec." % XLSX.split('/')[-1])
L.append("// FieldMapping rows: %d   RiskRules rows: %d" % (len(fields), len(rules)))
L.append("")
L.append("def round1(v) { v == null ? null : (Math.round((v as Double) * 10.0) / 10.0) }")
L.append("def n(v)      { v == null ? Double.NEGATIVE_INFINITY : (v as Double) }")
L.append("")
L.append("def Message processData(Message message) {")
L.append("    def src = new JsonSlurper().parseText(message.getBody(String) ?: '{}')")
L.append("    def t = [:]")
L.append("")
for f in fields:
    e = expr(f["Source Path"] or "", f["Transform"])
    if e:
        L.append(f"    t.{f['Target Field']} = {e}")
L.append("")
L.append("    // integration-test overrides (properties set by the sender hook)")
L.append("    ['tempF':'simTempF','aqi':'simAqi','windMph':'simWindMph'].each { k, prop ->")
L.append("        def o = message.getProperty(prop)")
L.append("        if (o != null) { t[k] = (o as String).isNumber() ? (o as Double) : t[k] }")
L.append("    }")
L.append("")
L.append("    // --- decision table from sheet 'RiskRules', evaluated by Priority ---")
first = True
for r in rules:
    kw = "if" if first else "} else if"; first = False
    L.append(f"    {kw} ({cond(r['Condition'])}) {{        // {r['Rule ID']} prio {r['Priority']}")
    L.append(f"        t.riskLevel = '{r['riskLevel']}'")
    L.append(f"        t.advisory  = '{r['advisory']}'")
    L.append(f"        t.dispatchFlag = {str(r['dispatchFlag']).strip().upper() == 'TRUE'}".lower().replace("t.dispatchflag","t.dispatchFlag"))
    L.append(f"        t.ruleApplied = '{r['Rule ID']}'")
L.append("    }")
L.append("")
req = [f["Target Field"] for f in fields if str(f.get("Required","")).upper() == "Y"]
L.append("    def missing = %s.findAll { t[it] == null }" % ("['" + "','".join(req) + "']"))
L.append("    if (missing) {")
L.append("        message.setProperty('errorCode','MAPPING_INCOMPLETE')")
L.append("        throw new RuntimeException('MAPPING_INCOMPLETE: ' + missing.join(','))")
L.append("    }")
L.append("")
L.append("    message.setBody(JsonOutput.toJson([site: t, sourcePayload: src]))")
L.append("    message.setHeader('Content-Type','application/json')")
L.append("    return message")
L.append("}")

open(OUT, "w").write("\n".join(L) + "\n")
print(f"generated {OUT}: {len(fields)} field mappings, {len(rules)} conditional rules -> {len(L)} lines")
